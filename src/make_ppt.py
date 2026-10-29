"""Run from the project root:  python src/make_ppt.py
Builds presentation/project_slides.pptx with text, figures and tables."""
import os
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)
blank = prs.slide_layouts[6]


def slide(title):
    s = prs.slides.add_slide(blank)
    tb = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.9))
    tb.text_frame.text = title
    p = tb.text_frame.paragraphs[0]
    p.font.size, p.font.bold = Pt(32), True
    return s


def bullets(s, lines, left=0.6, top=1.4, width=12, height=5.5, size=20):
    tf = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height)).text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text, p.font.size = "\u2022 " + line, Pt(size)
        p.space_after = Pt(8)


def pic(s, name, left, top, width):
    path = "figures/" + name
    if os.path.exists(path):
        s.shapes.add_picture(path, Inches(left), Inches(top), width=Inches(width))
    else:
        bullets(s, [f"[missing figure: {path}]"], left, top, width, 1)


def table(s, df, left=0.6, top=1.4, width=12, size=12):
    df = df.reset_index()
    rows, cols = df.shape[0] + 1, df.shape[1]
    t = s.shapes.add_table(rows, cols, Inches(left), Inches(top), Inches(width),
                           Inches(0.35 * rows)).table
    for j, c in enumerate(df.columns):
        t.cell(0, j).text = str(c)
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(row):
            t.cell(i, j).text = f"{v:.4f}" if isinstance(v, float) else str(v)
    for r in t.rows:
        for c in r.cells:
            for p in c.text_frame.paragraphs:
                p.font.size = Pt(size)


s = slide("Stock Price Prediction using Supervised Machine Learning")
bullets(s, ["Predicting Microsoft's next-week return with five ML frameworks",
            "Team: [names]", "[Course / date]"], size=24)

s = slide("Problem and objective")
bullets(s, ["Target: MSFT weekly return, one week ahead",
            "Why returns: prices are non-stationary (ADF p = 0.993), returns are stationary",
            "Features: correlated assets (IBM, GOOGL, USD/JPY, GBP/USD, S&P 500, Dow, VIX)",
            "Compare XGBoost, Keras, fastai, TabNet, Lightning vs baselines",
            "Criteria: RMSE, MAE, R\u00b2, R\u00b2_OOS, direction; simplicity, transparency, efficacy"])

s = slide("Data and features")
bullets(s, ["Yahoo Finance, daily closes 2010 to Oct 2026, 8 series",
            "Aligned to MSFT trading days; forward-fill only (never back-fill)",
            "876 weekly returns; current + 4 lagged weeks = 40 features, 871 samples",
            "Extension: 5 technical indicators (MA ratios, EMA ratio, momentum, volatility)"])

s = slide("Methodology and leakage controls")
bullets(s, ["Target shifted one week: features known at Friday close",
            "Chronological 80/20 split: train 696 weeks, test 175 weeks (Jun 2023 to Oct 2026)",
            "Scaler fit on training data only; TimeSeriesSplit for tuning",
            "Neural nets: last 15% of training as validation, early stopping, 3 seeds averaged",
            "Colab for Keras, fastai, TabNet, Lightning"])

s = slide("EDA: returns")
pic(s, "02_msft_return_dist.png", 0.4, 1.3, 6.3)
pic(s, "03_corr_heatmap.png", 6.9, 1.3, 6.0)
bullets(s, ["Fat tails (excess kurtosis 3.6, Jarque-Bera p ~ 4e-104); same-week correlation "
            "with S&P 500 is 0.67 but cannot be used for prediction"], 0.6, 6.2, 12, 1.2, 16)

s = slide("EDA: predictability")
pic(s, "06_acf_pacf.png", 0.4, 1.3, 7.0)
pic(s, "05_rolling_vol.png", 7.6, 1.3, 5.3)
bullets(s, ["No autocorrelation in returns; volatility clusters. Best lagged feature |r| = 0.099"],
        0.6, 6.4, 12, 1, 16)

s = slide("Models used")
bullets(s, ["Baselines: train mean, zero, linear regression, Ridge (alpha ~ 2,976)",
            "XGBoost: default vs tuned (depth 3, 100 trees, lr 0.02, L2 = 50)",
            "Keras MLP (32-16), fastai tabular, PyTorch TabNet, PyTorch Lightning MLP",
            "Flash failed to install (needs numpy<1.24): plain PyTorch Lightning used instead"])

s = slide("Results: base features")
table(s, pd.read_csv("results/final_comparison_base.csv", index_col=0))

s = slide("Model comparison")
pic(s, "19_model_comparison.png", 0.4, 1.3, 12.5)

s = slide("Key findings")
bullets(s, ["RMSE spread of only 0.0007 (excl. default XGBoost): no clear winner",
            "Best R\u00b2_OOS is 0.023; direction accuracy ~ 52% baseline, not beaten significantly",
            "Default XGBoost overfits (R\u00b2 -0.08); regularization fixes it",
            "Technical indicators did not help consistently; TabNet flipped sign (unstable)",
            "Training loss falls, validation loss flat: little signal to learn"])

s = slide("Simplicity / Transparency / Efficacy")
rows = [("Linear regression", 5, 5, 3), ("Ridge", 5, 5, 3), ("XGBoost (tuned)", 4, 4, 3),
        ("Keras MLP", 3, 2, 3), ("fastai tabular", 3, 2, 3), ("TabNet", 2, 3, 2),
        ("PyTorch Lightning", 2, 2, 3)]
table(s, pd.DataFrame(rows, columns=["model", "Simplicity", "Transparency", "Efficacy"])
      .set_index("model").astype(str))
bullets(s, ["Scores 1 to 5 (5 = best); efficacy differences are within noise"], 0.6, 5.3, 12, 1, 16)

s = slide("Expected returns")
pic(s, "20_expected_returns.png", 0.4, 1.3, 7.8)
bullets(s, ["Forecasts sit near historical means (MSFT 0.43%, IBM 0.21%, GOOGL 0.44% per week)",
            "IBM and GOOGL: R\u00b2_OOS \u2264 0 for all frameworks",
            "Models disagree widely: use the historical mean as the anchor",
            "Not investment advice"], 8.4, 1.4, 4.6, 5, 17)

s = slide("Limitations, conclusion, future work")
bullets(s, ["Weekly MSFT returns are almost unpredictable from these features",
            "Limits: single split, 2026 spikes in test set, no fundamentals/news, small samples",
            "Future: walk-forward validation, news sentiment, fundamentals, volatility or "
            "direction targets, longer horizons"])

os.makedirs("presentation", exist_ok=True)
prs.save("presentation/project_slides.pptx")
print("saved presentation/project_slides.pptx")
